"""Carga de xlsx: happy path openpyxl + fallback ante .rels rotos."""
from __future__ import annotations

import io
import unittest
import zipfile
from datetime import date, datetime
from unittest.mock import patch
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook

from excel_io import load_dossier_workbook, workbook_from_tabular_bytes


def _dossier_xlsx_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "NoticiaId",
            "Fecha",
            "Medio",
            "Tipo de Medio",
            "Título",
            "Link Nota",
            "Empresa rel.",
        ]
    )
    ws.append(
        [
            101,
            date(2026, 1, 15),
            "El Tiempo",
            "online",
            "Nissan lanza modelo",
            "ver nota",
            "Nissan",
        ]
    )
    ws.cell(row=2, column=6).hyperlink = "https://example.com/nota"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _corrupt_worksheet_rels(xlsx_bytes: bytes, payload: bytes | None = None) -> bytes:
    """Inyecta XML inválido en xl/worksheets/_rels/*.rels (el crash de get_dependents)."""
    in_buf = io.BytesIO(xlsx_bytes)
    out_buf = io.BytesIO()
    with zipfile.ZipFile(in_buf, "r") as zin, zipfile.ZipFile(out_buf, "w") as zout:
        found = False
        for info in zin.infolist():
            content = zin.read(info.filename)
            if "worksheets/_rels/" in info.filename.replace("\\", "/") and info.filename.endswith(
                ".rels"
            ):
                found = True
                if payload is not None:
                    content = payload
                else:
                    text = content.decode("utf-8")
                    # Target con & sin escapar → ParseError en fromstring
                    broken = (
                        '<Relationship Id="rIdBad" '
                        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
                        'Target="https://example.com/a?b=1&c=2" TargetMode="External"/>'
                    )
                    if "</Relationships>" in text:
                        text = text.replace("</Relationships>", broken + "</Relationships>")
                    else:
                        text = broken
                    content = text.encode("utf-8")
            zout.writestr(info.filename, content)
        if not found:
            # openpyxl a veces no escribe sheet rels; forzar el archivo que get_dependents lee
            zout.writestr(
                "xl/worksheets/_rels/sheet1.xml.rels",
                payload
                or (
                    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                    b'<Relationship Id="rIdBad" Type="http://example.com/rel" '
                    b'Target="https://example.com/a?b=1&c=2"/></Relationships>'
                ),
            )
    return out_buf.getvalue()


class LoadDossierWorkbookTests(unittest.TestCase):
    def test_happy_path_openpyxl_file_still_loads(self):
        data = _dossier_xlsx_bytes()
        wb = load_dossier_workbook(io.BytesIO(data), data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "NoticiaId")
        self.assertEqual(rows[1][0], 101)
        self.assertEqual(rows[1][4], "Nissan lanza modelo")
        link_cell = wb.active.cell(row=2, column=6)
        self.assertEqual(link_cell.value, "ver nota")
        self.assertTrue(
            link_cell.hyperlink and "example.com/nota" in str(link_cell.hyperlink.target)
        )

    def test_broken_rels_crashes_openpyxl_but_fallback_loads(self):
        broken = _corrupt_worksheet_rels(_dossier_xlsx_bytes())

        with self.assertRaises(ParseError):
            load_workbook(io.BytesIO(broken), data_only=True)

        wb = load_dossier_workbook(io.BytesIO(broken), data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(rows[0][5], "Link Nota")
        self.assertEqual(rows[1][4], "Nissan lanza modelo")
        self.assertEqual(rows[1][6], "Nissan")
        self.assertEqual(wb.active.cell(row=2, column=6).value, "ver nota")

    def test_garbage_rels_still_loads_via_sanitize_or_calamine(self):
        broken = _corrupt_worksheet_rels(
            _dossier_xlsx_bytes(),
            payload=b"this is not xml &&& <Relationships",
        )

        with self.assertRaises(Exception):
            load_workbook(io.BytesIO(broken), data_only=True)

        wb = load_dossier_workbook(io.BytesIO(broken), data_only=True)
        titles = [c.value for c in wb.active[1]]
        self.assertIn("Título", titles)

    def test_calamine_rebuild_when_openpyxl_always_fails(self):
        data = _dossier_xlsx_bytes()

        def boom(*_args, **_kwargs):
            raise ParseError("not well-formed (invalid token)")

        with patch("excel_io.load_workbook", side_effect=boom):
            wb = load_dossier_workbook(io.BytesIO(data), data_only=True)

        rows = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "NoticiaId")
        self.assertEqual(rows[1][4], "Nissan lanza modelo")

    def test_spanish_error_when_bytes_are_not_xlsx(self):
        with self.assertRaises(RuntimeError) as ctx:
            load_dossier_workbook(io.BytesIO(b"esto no es un excel"), data_only=True)
        self.assertIn("No se pudo leer el archivo Excel", str(ctx.exception))

    def test_workbook_from_tabular_bytes_roundtrip(self):
        wb = workbook_from_tabular_bytes(_dossier_xlsx_bytes())
        values = [c.value for c in next(wb.active.iter_rows(min_row=2, max_row=2))]
        self.assertEqual(values[3], "online")
        self.assertTrue(
            values[1] == date(2026, 1, 15)
            or isinstance(values[1], datetime)
            or str(values[1]).startswith("2026-01-15")
        )


if __name__ == "__main__":
    unittest.main()
